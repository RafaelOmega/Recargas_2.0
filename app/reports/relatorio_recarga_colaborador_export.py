# -*- coding: utf-8 -*-
"""Exportação do relatório de recargas por colaborador para XLSX e CSV.

Mesmos dados e mesma ordem do PDF. XLSX grava valores numéricos com
formato de célula; CSV grava valores já formatados no padrão brasileiro
(separador ';' e BOM para o Excel).
"""
import csv
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models.relatorio_recarga_colaborador import ItemRecargaRelatorio
from app.reports.relatorio_recarga_colaborador_pdf import (
    _formatar_documento, _moeda,
)
from app.utils.logger import get_logger

logger = get_logger("relatorio_recarga_colaborador_export")

AZUL = "1F3B5B"
ZEBRA = "F2F5F8"
TOTAL = "E8EDF2"

_CABECALHOS = ["Data", "Seq", "Colaborador", "CPF", "Carga",
               "Sem Carga", "Valor", "Benefício"]
_FMT_MOEDA = '"R$" #,##0.00'


# ---------------- XLSX ----------------

def gerar_xlsx_recarga_colaborador(registros: list[ItemRecargaRelatorio],
                                   caminho: str,
                                   subtitulo: str = "") -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Recargas por Colaborador"

    ws["A1"] = "RELATÓRIO DE RECARGAS POR COLABORADOR"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    linha = 2
    if subtitulo:
        ws.cell(linha, 1, subtitulo)
        linha += 1
    ws.cell(linha, 1, f"Emitido em {datetime.now():%d/%m/%Y %H:%M}")
    linha += 2

    for coluna, texto in enumerate(_CABECALHOS, start=1):
        cel = ws.cell(linha, coluna, texto)
        cel.font = Font(bold=True, color=AZUL)
        cel.fill = PatternFill("solid", fgColor=ZEBRA)
    linha += 1

    for indice, r in enumerate(registros):
        fundo = ZEBRA if indice % 2 else "FFFFFF"
        ws.cell(linha, 1, r.data.strftime("%d/%m/%Y") if r.data else "")
        ws.cell(linha, 2, r.sequencia)
        ws.cell(linha, 3, r.nome)
        ws.cell(linha, 4, _formatar_documento(r.documento))
        ws.cell(linha, 5, r.carga)
        ws.cell(linha, 6, r.sem_carga)
        cel_valor = ws.cell(linha, 7, r.valor)
        cel_valor.number_format = _FMT_MOEDA
        ws.cell(linha, 8, r.beneficio)
        for coluna in range(1, 9):
            ws.cell(linha, coluna).fill = PatternFill("solid",
                                                      fgColor=fundo)
        linha += 1

    total = sum(r.valor for r in registros)
    cel_rotulo = ws.cell(linha, 6, f"Totais — {len(registros)} item(ns)")
    cel_rotulo.font = Font(bold=True)
    cel_rotulo.fill = PatternFill("solid", fgColor=TOTAL)
    cel_total = ws.cell(linha, 7, total)
    cel_total.font = Font(bold=True)
    cel_total.fill = PatternFill("solid", fgColor=TOTAL)
    cel_total.number_format = _FMT_MOEDA

    for coluna, largura in zip("ABCDEFGH",
                               (12, 8, 30, 16, 12, 12, 14, 14)):
        ws.column_dimensions[coluna].width = largura

    wb.save(caminho)
    logger.info("XLSX gerado: %s", caminho)
    return caminho


# ---------------- CSV ----------------

def gerar_csv_recarga_colaborador(registros: list[ItemRecargaRelatorio],
                                  caminho: str,
                                  subtitulo: str = "") -> str:
    with open(caminho, "w", newline="", encoding="utf-8-sig") as arquivo:
        writer = csv.writer(arquivo, delimiter=";")
        writer.writerow(["RELATÓRIO DE RECARGAS POR COLABORADOR"])
        if subtitulo:
            writer.writerow([subtitulo])
        writer.writerow([f"Emitido em {datetime.now():%d/%m/%Y %H:%M}"])
        writer.writerow(_CABECALHOS)
        for r in registros:
            writer.writerow([
                r.data.strftime("%d/%m/%Y") if r.data else "",
                r.sequencia,
                r.nome,
                _formatar_documento(r.documento),
                r.carga,
                r.sem_carga,
                _moeda(r.valor),
                r.beneficio,
            ])
        total = sum(r.valor for r in registros)
        writer.writerow([
            "", "", "", "", "",
            f"Totais — {len(registros)} item(ns)",
            _moeda(total), "",
        ])
    logger.info("CSV gerado: %s", caminho)
    return caminho
